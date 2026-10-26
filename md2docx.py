#!/usr/bin/env python3
"""Minimal markdown -> docx for ABF reports: headings, bold (**), bullets (-), pipe tables."""
import sys, re
from docx import Document
from docx.shared import Pt

md_path, out_path = sys.argv[1], sys.argv[2]
doc = Document()
lines = open(md_path).read().split('\n')

def add_runs(par, text):
    # split on **bold**
    for i, seg in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if seg == '': continue
        r = par.add_run(seg)
        if i % 2 == 1:
            r.bold = True

i = 0
while i < len(lines):
    line = lines[i].rstrip('\n')
    # table block
    if line.strip().startswith('|') and i+1 < len(lines) and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i+1]):
        header = [c.strip() for c in line.strip().strip('|').split('|')]
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        t = doc.add_table(rows=1, cols=len(header))
        t.style = 'Light Grid Accent 1'
        for j, h in enumerate(header):
            add_runs(t.rows[0].cells[j].paragraphs[0], h)
        for row in rows:
            cells = t.add_row().cells
            for j in range(len(header)):
                add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else '')
        doc.add_paragraph('')
        continue
    if line.startswith('# '):
        doc.add_heading(line[2:], level=0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=2)
    elif line.startswith('> '):
        p = doc.add_paragraph(); p.style = 'Intense Quote'; add_runs(p, line[2:])
    elif re.match(r'^\s*[-*]\s+', line):
        indent = len(line) - len(line.lstrip())
        p = doc.add_paragraph(style='List Bullet' if indent < 2 else 'List Bullet 2')
        add_runs(p, re.sub(r'^\s*[-*]\s+', '', line))
    elif line.strip() == '':
        pass
    else:
        p = doc.add_paragraph(); add_runs(p, line)
    i += 1

doc.save(out_path)
print('wrote', out_path)
