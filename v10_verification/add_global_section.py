"""Append a 'Global numbers' section (two combined cross-platform tables, ITT) to the OTT v10 doc.
Binoop-style weighting: combined lift = 0.85*OTT + 0.10*mobile + 0.05*web; CI treats platforms as independent.
Usage: add_global_section.py <doc.docx> <results_dir>"""
import json, math, os, sys, re
from copy import deepcopy
import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DOC, RD = sys.argv[1:3]
W = {"ott": .85, "mobile": .10, "web": .05}
V = {p: json.load(open(os.path.join(RD, f"values_{p}.json")))["head"]["ITT"] for p in W}
ARMS = ("v4", "v5b", "v6")
s = lambda x: f"{x:+.3f}".replace("-", "−")

def cell(m, a):
    l = sum(W[p] * V[p][m][a]["l"] for p in W)
    se = math.sqrt(sum((W[p] * (V[p][m][a]["hi"] - V[p][m][a]["lo"]) / 2 / 1.96) ** 2 for p in W))
    lo, hi = l - 1.96 * se, l + 1.96 * se
    return f"{s(l)}% [{s(lo)}, {s(hi)}] {'sig' if lo > 0 or hi < 0 else 'ns'}"

d = docx.Document(DOC)
if any(p.text.strip() == "Global numbers" for p in d.paragraphs):
    sys.exit("already present")
src_tbl = d.tables[2]._tbl

def settext(tc, text):
    p = tc.findall(qn("w:p"))[0]
    runs = p.findall(qn("w:r"))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn("w:t")).text = text

def make_table(rev, tvt):
    t = deepcopy(src_tbl)
    rows = t.findall(qn("w:tr"))
    for r in rows[4:]:
        t.remove(r)
    grid = t.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    t.find(qn("w:tblGrid")).remove(grid[4])
    hdr = ["Population", "Arm", "n", "Unbudgeted revenue", "TVT"]
    for ri, r in enumerate(t.findall(qn("w:tr"))):
        tcs = r.findall(qn("w:tc"))
        r.remove(tcs[4])          # drop Impressions
        tcs = r.findall(qn("w:tc"))
        if ri == 0:
            vals = hdr
        else:
            a = ARMS[ri - 1]
            n = sum(V[p]["n"][a] for p in W)
            vals = ["ITT", a, f"{n:,}", cell(rev, a), cell(tvt, a)]
        for tc, v in zip(tcs, vals):
            settext(tc, v)
    return t

h = d.add_paragraph("Global numbers", style="Heading 1")
intro = d.add_paragraph(
    "Cross-platform view of the global ITT lifts versus ABF2 Control (unbudgeted revenue and TVT; CUPED, 95% intervals). Method follows the independent OTT report’s Appendix C: "
    "combined lift = 0.85 × OTT lift + 0.10 × mobile lift + 0.05 × web lift (shares of company revenue and TVT, supplied by the business and applied to both metrics); "
    "the interval treats the three platforms as independent (separate device populations). Window 2026-10-09 → 2026-10-19, pre-period 2026-09-25 → 2026-10-08. "
    "n is the sum of ITT devices per arm across the three platforms, shown for reference (the weights are not device shares). "
    "It assumes each experiment’s global ITT lift approximates the effect of a full rollout on that platform. Sig = interval excludes zero (nominal 95%, not adjusted for sequential looks or three arms).")
c1 = d.add_paragraph("Combined global unbudgeted detail (all titles):")
c1.runs[0].bold = True
anchor = c1._p
t1 = make_table("revu", "tvt"); anchor.addnext(t1)
c2 = d.add_paragraph("Combined on-title (the 224,739 targeted titles):")
c2.runs[0].bold = True
t1.addnext(c2._p)
t2 = make_table("orevu", "otvt"); c2._p.addnext(t2)
note = d.add_paragraph(
    "Note on the on-title table: the 85/10/5 weights are company-level revenue and TVT shares, so applying them to on-title lifts is an approximation — each on-title lift is measured against that platform’s targeted-title baseline, "
    "and targeted titles may be a different share of each platform’s revenue. Per-platform lifts are in the OTT tables above and in the Mobile and Web reports (same window and title set).")
t2.addnext(note._p)
# the add_paragraph calls appended intro/c1/c2/note at the end; fix order: heading, intro, c1, t1, c2, t2, note
body = d.element.body
sect = body.find(qn("w:sectPr"))
for el in (h._p, intro._p, c1._p, t1, c2._p, t2, note._p):
    body.remove(el)
for el in (h._p, intro._p, c1._p, t1, c2._p, t2, note._p):
    sect.addprevious(el)
d.save(DOC)
print("ok")
