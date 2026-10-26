"""Refresh the ABF3 v10 OTT results doc (in place, formatting preserved) with the 2026-10-20 re-run.

Usage: edit_ott_r3.py <src.docx> <results_dir> <prev_dir> <out.docx>
  src = the original 10-15 doc (template), results_dir = results_20261020, prev_dir = prev_20261019/results (previous published numbers).
Tables are overwritten from the results files; prose paragraphs that quote numbers are rewritten from the same numbers
(matched by their opening words). §B.5 pacing / pre-roll analysis is left as-is and flagged as the 10-15 read.
Window 2026-10-09 -> end of 2026-10-19. On-title/ToT/bands/type/below use the 224,739 targeted titles
(ml_dev.tmp.abf3_v10_targeted_titles), not the 299k live ABF2 catalog.
"""
import json, os, re, sys
from copy import deepcopy
import docx
from docx.text.paragraph import Paragraph

SRC, RD, PD, OUT = sys.argv[1:5]
J = lambda d, f: json.load(open(os.path.join(d, f)))
V = J(RD, "values_ott.json"); H, R = V["head"], V["raw"]
OLD = J(PD, "values_ott.json")["head"]
OLD_R = J(PD, "values_ott.json")["raw"]
OLD15 = J(os.path.join(os.path.dirname(os.path.abspath(__file__)), "prev_20261015"), "values.json")["head"]
G = J(RD, "gate_ott_base.json")
SV = J(RD, "served.json")
d = docx.Document(SRC)
ARMS = ("v4", "v5b", "v6")
BANDS = ("1-3", "4-6", "7-9", "10-12", "13+")
MISSES = []


# ---------- formatting ----------
def s(x, nd):
    x = round(x, nd) + 0.0
    return f"{x:+.{nd}f}".replace("-", "−")


sig = lambda e: "sig" if e["lo"] > 0 or e["hi"] < 0 else "ns"
issig = lambda e: e["lo"] > 0 or e["hi"] < 0


def full(e, nd=3, tag=True):
    t = f"{s(e['l'], nd)}% [{s(e['lo'], nd)}, {s(e['hi'], nd)}]"
    return t + (" " + sig(e) if tag else "")


def short(e, nd=2):
    return f"{s(e['l'], nd)}% {sig(e)}"


def tvt_cell(e, nd=2):
    return full(e, nd) if sig(e) == "sig" else short(e, nd)


def pc(e, nd=3):          # "+0.853%"  (no tag)
    return f"{s(e['l'], nd)}%"


def pt(e, nd=3):          # "+0.853% sig"
    return f"{s(e['l'], nd)}% {sig(e)}"


n = lambda x: f"{x:,}"
band_lbl = lambda b: b.replace("-", "–")
I = H["ITT"]; IO = OLD["ITT"]
ctrl_n = I["n"]["Control"]


# ---------- docx helpers ----------
def setcell(cell, text):
    runs = [r for p in cell.paragraphs for r in p.runs]
    if not runs:
        cell.paragraphs[0].add_run(text); return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def set_markup(p, markup):
    """Replace the paragraph's runs with runs parsed from **bold** markup."""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for i, part in enumerate(re.split(r"\*\*", markup)):
        if part == "":
            continue
        run = p.add_run(part)
        if i % 2 == 1:
            run.bold = True


def find(prefix, nth=0):
    hits = [p for p in d.paragraphs if p.text.startswith(prefix)]
    if len(hits) <= nth:
        MISSES.append(prefix); return None
    return hits[nth]


def setp(prefix, markup, nth=0):
    p = find(prefix, nth)
    if p is not None:
        set_markup(p, markup)
    return p


def rep(old, new):
    hits = 0
    pars = list(d.paragraphs) + [p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for p in pars:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); hits += 1
    if hits == 0:
        MISSES.append(old)


def insert_after(p, markup, template=None):
    new = deepcopy((template or p)._p)
    p._p.addnext(new)
    np_ = Paragraph(new, p._parent)
    set_markup(np_, markup)
    return np_


# ---------- tables ----------
T = d.tables
for g in ARMS:
    row0 = {"v4": 1, "v6": 2, "v5b": 3}[g]
    setcell(T[0].rows[row0].cells[1], full(I["revt"][g]))
    setcell(T[0].rows[row0].cells[2], short(I["revb"][g], 3))
    setcell(T[0].rows[row0].cells[3], short(I["revu"][g], 3))
    setcell(T[0].rows[row0].cells[4], short(I["impt"][g], 3))
    setcell(T[0].rows[row0].cells[5], short(I["tvt"][g], 3))
    row1 = {"v4": 1, "v5b": 2, "v6": 3}[g]
    for ci, m in enumerate(("revt", "revb", "revu", "impt"), start=1):
        setcell(T[1].rows[row1].cells[ci], full(I[m][g]))
for bi, pop in enumerate(("ITT", "Active", "ToT")):
    for ai, g in enumerate(ARMS):
        r = 1 + 3 * bi + ai; P = H[pop]
        setcell(T[2].rows[r].cells[2], n(P["n"][g]))
        setcell(T[2].rows[r].cells[3], full(P["revu"][g]))
        setcell(T[2].rows[r].cells[4], full(P["impu"][g]))
        setcell(T[2].rows[r].cells[5], full(P["tvt"][g]))
        setcell(T[3].rows[r].cells[2], full(P["orevu"][g]))
        setcell(T[3].rows[r].cells[3], full(P["oimpu"][g]))
        setcell(T[3].rows[r].cells[4], full(P["otvt"][g]))
for bi, pop in enumerate(("NPV", "PV")):
    for ai, g in enumerate(ARMS):
        r = 1 + 3 * bi + ai; e = H[pop]["tvt"][g]
        setcell(T[4].rows[r].cells[2], n(H[pop]["n"][g]))
        setcell(T[4].rows[r].cells[3], full(e, 3, tag=False))
        setcell(T[4].rows[r].cells[4], "Yes" if sig(e) == "sig" else "No")
for bi, b in enumerate(BANDS):
    for ai, g in enumerate(ARMS):
        r = 1 + 3 * bi + ai; x = R["band|" + b][g]
        setcell(T[5].rows[r].cells[2], full(x["rev"], 2))
        setcell(T[5].rows[r].cells[3], short(x["tvt"]))
        setcell(T[5].rows[r].cells[4], full(x["imps"], 2))
        setcell(T[5].rows[r].cells[5], n(x["n"]))
r = 1
for b in ("10-12", "13+"):
    for tag in ("below", "at/above"):
        for g in ARMS:
            x = R[f"below|{g}|{b}|{tag}"][g]
            setcell(T[6].rows[r].cells[3], full(x["rev"], 2))
            setcell(T[6].rows[r].cells[4], tvt_cell(x["tvt"]))
            r += 1
for ti, tt in enumerate(("Movie", "Episode")):
    for ai, g in enumerate(ARMS):
        r = 1 + 3 * ti + ai; x = R["type|" + tt][g]
        setcell(T[7].rows[r].cells[2], full(x["rev"], 2))
        setcell(T[7].rows[r].cells[3], short(x["tvt"]))
        setcell(T[7].rows[r].cells[4], short(x["imps"]))
        setcell(T[7].rows[r].cells[5], n(x["n"]))
# served-list tables (static, re-checked)
cp = {(r[1], r[2]): r for r in SV["rows"] if r[0] == "cphr"}
mins = lambda x: f"~{float(x):.0f}s ({float(x)/60:.1f} min)"
setcell(T[8].rows[1].cells[1], mins(cp[("Episode", "v4")][8])); setcell(T[8].rows[1].cells[2], mins(cp[("Movie", "v4")][8]))
for g in ARMS:
    row = [rw for rw in T[8].rows if rw.cells[0].text.strip() == g][0]
    setcell(row.cells[1], mins(cp[("Episode", g)][9])); setcell(row.cells[2], mins(cp[("Movie", g)][9]))
cell = {(c["tt"], c["band"], c["arm"]): c["sim_delta"] for c in G["cells"]}
for row in T[9].rows[1:]:
    tt, b = row.cells[0].text.strip(), row.cells[1].text.strip().replace("–", "-")
    for ci, g in enumerate(ARMS, start=2):
        setcell(row.cells[ci], f"{s(cell[(tt, b, g)], 1)}%")
dec = {x["arm"]: x for x in G["episode_decomp"]}
for row in T[10].rows[1:]:
    g = row.cells[0].text.strip()
    for ci, k in enumerate(("total", "from_fewer", "from_placement", "from_gains"), start=1):
        setcell(row.cells[ci], f"{s(dec[g][k], 1)}%")
for row in T[11].rows[1:]:
    tt = row.cells[0].text.strip(); base = float(cp[(tt, "v4")][6])
    setcell(row.cells[1], f"{base:.2f}")
    for ci, g in enumerate(ARMS, start=2):
        setcell(row.cells[ci], f"{float(cp[(tt, g)][7]):.2f} ({s((float(cp[(tt, g)][7]) / base - 1) * 100, 0)}%)")
bd = {(r[1], r[4], r[2]): r for r in SV["rows"] if r[0] == "band"}
for row in T[12].rows[1:]:
    tt, b, g = row.cells[0].text.strip(), row.cells[1].text.strip().replace("–", "-"), row.cells[2].text.strip()
    for ci, k in enumerate((10, 11, 12), start=3):
        setcell(row.cells[ci], bd[(tt, b, g)][k])
for ai, g in enumerate(ARMS):
    setcell(T[13].rows[1 + ai].cells[1], short(I["impb"][g], 3))
    setcell(T[13].rows[1 + ai].cells[2], short(I["impu"][g], 3))

# ---------- prose ----------
oldn = IO["n"]["Control"]
chg = []
for m, lbl in (("revt", "total revenue"), ("revb", "budgeted revenue"), ("revu", "unbudgeted revenue"), ("impt", "impressions"), ("tvt", "TVT")):
    for a in ARMS:
        if issig(IO[m][a]) != issig(I[m][a]):
            chg.append(f"{a} {lbl} {'ns → sig' if issig(I[m][a]) else 'sig → ns'} ({s(IO[m][a]['l'],3)}% → {s(I[m][a]['l'],3)}%)")
move = "; ".join(f"{a} {s(IO['revt'][a]['l'],3)}% → {s(I['revt'][a]['l'],3)}%" for a in ARMS)
tmove = "; ".join(f"{a} {s(IO['tvt'][a]['l'],3)}% → {s(I['tvt'][a]['l'],3)}%" for a in ARMS)
OO = IO["orevu"]
oto = lambda a: f"{pc(OO[a],3)} → {pc(I['orevu'][a],3)}"
refresh = (f"Window extended from the end of 2026-10-18 to the end of 2026-10-19 (every complete day loaded in the exposures table); ITT devices/arm {oldn/1e6:.2f}M → {ctrl_n/1e6:.2f}M. "
           "**Title set corrected:** the on-title, ToT, band, type and below-control tables now use the 224,739 targeted titles (`ml_dev.tmp.abf3_v10_targeted_titles`, the titles served v10 cue points) instead of the 299k live ABF2 catalog, "
           "which included about 74.5k titles that carry identical cues in every arm and so diluted every on-title and band lift. "
           f"On-title ITT unbudgeted revenue, previous version → now (title set and window both changed): v4 {oto('v4')}; v5b {oto('v5b')}; v6 {oto('v6')}. "
           f"Global ITT total revenue, previous version → now: {move}. TVT: {tmove}. "
           + ("Significance calls that changed: " + "; ".join(chg) + "." if chg else "No global ITT significance call changed.")
           + " The corrected on-title and global numbers are consistent with an independent OTT analysis of the same experiment (a slightly different window). Every table below except §B.5's pacing analysis was re-run.")

setp("Window:", "**Window:** **2026-10-09 → 2026-10-19 (complete days)** — outcomes are capped at 2026-10-20 00:00 UTC; the exposures table is loaded through the end of 2026-10-19 (2026-10-20 is a partial day and is excluded). "
     "The first version of this report (to 2026-10-15 10:00 EDT) was matched to the Statsig scorecard; this one has not been re-reconciled with it. This is an early read of a live test, not a graduation read.")
setp("Method:", f"**Method:** device-level CUPED, pre-period 2026-09-25 → 2026-10-08, Welch 95% intervals. ~{ctrl_n/1e6:.2f}M devices/arm (ITT). Revenue is reported three ways — **total** (all billable ads, the platform topline), "
     "**budgeted** (campaigns that carry a budget: direct, guaranteed, PMP), and **unbudgeted** (open-auction only, the v8 report’s metric) — plus ad impressions and TVT. Definitions are in Appendix C. Report date 2026-10-20.")
insert_after(find("Method:"), f"**Refresh (2026-10-20):** {refresh}")

tot_ratio = I["revt"]["v5b"]["l"] / I["revu"]["v5b"]["l"]
setp("v5b is the money arm.", f"**v5b is the money arm. It lifts total platform revenue about {I['revt']['v5b']['l']:.1f}% with a small TVT dip ({pc(I['tvt']['v5b'],2)}).** ~{ctrl_n/1e6:.2f}M devices per arm. "
     "The earlier read reconciled with the Statsig scorecard (Appendix C); that check has not been re-run on this window.")
setp("Read revenue on the total", "**Read revenue on the total, not just the open auction.** All three pruners serve more ads; the new ads are filled mostly by **budgeted / guaranteed** demand, not the open auction. "
     f"v5b’s open-auction (unbudgeted) gain is now significant ({pc(I['revu']['v5b'])}) but **total** revenue rises ~{tot_ratio:.0f}× as much ({pc(I['revt']['v5b'])}). Ranking on **total revenue: v5b > v6 > v4**.")
setp("v5b: total revenue", f"**v5b: total revenue {pt(I['revt']['v5b'])}, budgeted {pt(I['revb']['v5b'])}, all-ad impressions {pt(I['impt']['v5b'])}.** On the 224,739 targeted titles, unbudgeted revenue is {pt(I['orevu']['v5b'])} "
     f"(on-title tables previously used the wider ABF2 catalog, which diluted it to {pc(OO['v5b'])}). TVT is a small negative ({pc(I['tvt']['v5b'])} global, {pc(H['PV']['tvt']['v5b'])} on power viewers; non-power {pt(H['NPV']['tvt']['v5b'])}) — minor, but the one engagement item to watch since v5b adds the most ad load.")
setp("v6: total revenue", f"**v6: total revenue {pt(I['revt']['v6'])}, budgeted {pt(I['revb']['v6'])}, all-ad impressions {pt(I['impt']['v6'])}, TVT flat ({pt(I['tvt']['v6'])}).** On-title unbudgeted revenue {pt(I['orevu']['v6'])}; "
     f"global unbudgeted is now also significant ({pt(I['revu']['v6'])}). The safer middle arm.")
setp("v4: roughly parity", f"**v4: roughly parity globally, negative on the treated titles.** Global total revenue {pt(I['revt']['v4'])} and global unbudgeted {pt(I['revu']['v4'])} (budgeted {pt(I['revb']['v4'])}) — close to its break-even design — "
     f"but on the targeted titles themselves it loses (on-title unbudgeted {pt(I['orevu']['v4'])}). Weakest of the three; TVT {pt(I['tvt']['v4'])}.")
setp("TVT is nearly flat everywhere", f"**TVT is nearly flat everywhere.** With ~{ctrl_n/1e6:.1f}M devices/arm the only significant global TVT move is v5b ({pc(I['tvt']['v5b'],2)}); v4 ({pc(I['tvt']['v4'],2)}) and v6 ({pc(I['tvt']['v6'],2)}) are not significant "
     f"(v4’s significant {pc(OLD15['ITT']['tvt']['v4'],2)} on the 10-15 read did not hold as the window grew). No arm shows a material engagement cost; v5b’s small negative is the only mild watch-item.")
epb = lambda g: [float(bd[("Episode", b, g)][10]) for b in ("7-9", "10-12", "13+")]
lo_b, hi_b = min(min(epb(g)) for g in ARMS), max(max(epb(g)) for g in ARMS)
mv, ep = R["type|Movie"], R["type|Episode"]
setp("Movies carry the gains", f"**Movies carry the gains; episodes leak — and the leak is understood and fixable.** On movies every arm gains unbudgeted revenue (v4 {pc(mv['v4']['rev'],2)}, v5b {pc(mv['v5b']['rev'],2)}, v6 {pc(mv['v6']['rev'],2)}, TVT flat). "
     f"On episodes v4 ({pc(ep['v4']['rev'],2)}) and v6 ({pc(ep['v6']['rev'],2)}) lose and v5b is {pt(ep['v5b']['rev'],2)} (its episode impressions are {pt(ep['v5b']['imps'],2)}); TVT flat. "
     f"The served cue lists (`ml_dev.tmp.abf3_new_pruner_cuepoints`) and the gate replay show why: all three arms prune dense episodes below the ABF2 count on ~{lo_b:.0f}–{hi_b:.0f}% of titles, and v4/v6 also place the first episode ad ~3.5 min late. "
     "A per-title floor plus an earlier episode cue would recover it (§2b, §2c, §B.4).")
setp("All figures are global ITT CUPED lifts versus Control. The offline", "All figures are global ITT CUPED lifts versus Control. The offline projections (v4 ≈ parity, v6 ≈ +0.43% revenue, v5b ≈ +1.31% revenue) were modelled on the open auction; "
     f"the live total-revenue gains land mostly through budgeted backfill instead (v5b open-auction {pc(I['revu']['v5b'],2)} vs the +1.31% projection). Direction is stable across ITT, Active, and ToT. "
     "Significance calls are nominal 95%, not adjusted for sequential looks or three arms. The Statsig scorecard widened intervals ~1.8× on the 10-15 read; applying that here would make v5b’s TVT dip not significant while v5b’s and v6’s total-revenue gains stay significant "
     "(approximation, not re-checked against Statsig for this window).")
setp("1. Headline", "1. Headline (2026-10-09 → 2026-10-19)")
setp("All figures are CUPED lifts versus Control, 95% intervals.", "All figures are CUPED lifts versus Control, 95% intervals. Global is across every title an exposed device watched. On-title is the same devices, restricted to the 224,739 targeted titles — the ABF2-catalog titles that were served v10 cue points (`ml_dev.tmp.abf3_v10_targeted_titles`). ToT is defined on the same set. “Sig” means the interval excludes zero (nominal p < 0.05).")
setp("On-title (ABF2 catalog only", "On-title (the 224,739 targeted titles — ABF2 titles served v10 cue points; same device counts as the matching global row):")
cm = I["ctrl_mean"]
setp("Control per-device means", f"Control per-device means: total revenue {cm['revt']:.3f}, budgeted {cm['revb']:.3f}, all-ad impressions {cm['impt']:.1f}; ITT n/arm ≈ {ctrl_n/1e6:.2f}M. The pattern is the key read: v5b and v6 serve more ads (impressions up), "
     f"the extra inventory is bought mostly by **budgeted / guaranteed** demand (budgeted revenue {pc(I['revb']['v5b'])} on v5b), so **total** revenue rises well above the open-auction gain (v5b unbudgeted {pc(I['revu']['v5b'])} vs total {pc(I['revt']['v5b'])}). Reading only the unbudgeted column understates v5b and v6.")
setp("Control n: ITT", f"Control n: ITT {n(ctrl_n)} · Active {n(H['Active']['n']['Control'])} · ToT {n(H['ToT']['n']['Control'])}.")
oa = lambda m, a: [H[p][m][a] for p in ("ITT", "Active", "ToT")]
setp("On the treated ABF2 titles, both", f"On the targeted titles, both **v5b and v6 show significant positive unbudgeted revenue** (ITT {pc(I['orevu']['v5b'])} and {pc(I['orevu']['v6'])}), "
     f"and v4 is significantly negative ({pc(oa('orevu','v4')[0])} ITT, {pc(oa('orevu','v4')[1])} Active, {pc(oa('orevu','v4')[2])} ToT). v4’s on-title TVT is no longer significant ({pt(I['otvt']['v4'])}), "
     "and its revenue and impressions are negative in every population, so v4 remains the weak arm on the treated catalog.")
pvs = H["PV"]["n"]["Control"] / ctrl_n * 100
oldpv = OLD15["PV"]["n"]["Control"]
added_pv = (H["PV"]["n"]["Control"] - oldpv) / (ctrl_n - OLD15["ITT"]["n"]["Control"]) * 100
setp("NPV_TVT", "**NPV_TVT — TVT split by viewer type before first exposure.** Power viewer = ≥8 visit-days in last 28 days and >4 watch-hours/visit-day from the latest `device_okr_daily` snapshot on/before first exposure. "
     f"Everyone else (incl. devices absent from the table) is non-power. Flag set pre-exposure; power share balanced (~{pvs:.1f}% per arm — down from the 10-15 read because the devices added since are only ~{added_pv:.1f}% power viewers).")
setp("Control n: non-power", f"Control n: non-power {n(H['NPV']['n']['Control'])} · power {n(H['PV']['n']['Control'])}. The only significant move is v5b on power viewers ({pt(H['PV']['tvt']['v5b'])}); v5b non-power is {pt(H['NPV']['tvt']['v5b'])}. "
     f"v4’s non-power {pc(H['NPV']['tvt']['v4'],2)} is no longer significant (interval {full(H['NPV']['tvt']['v4'],2,tag=False).split(' ',1)[1]}). v5b’s power-viewer dip is consistent with it adding the most ad load and is the one engagement item to keep watching.")
setp("Control n by band", "Control n by band: " + " · ".join(f"{band_lbl(b)} {n(R['band|'+b]['v4']['ctrl_n'])}" for b in BANDS) + ".")
b = lambda band, a, m="rev": R["band|" + band][a][m]
setp("10–12 is the positive pocket", f"**10–12 is the positive pocket for every arm**, biggest on v6 ({pc(b('10-12','v6'),2)}) then v5b ({pc(b('10-12','v5b'),2)}); v4 {pc(b('10-12','v4'),2)}. TVT flat.")
setp("13+ loses on", f"**13+** loses on v4 and v5b ({pc(b('13+','v4'),2)}, {pc(b('13+','v5b'),2)}); v6 is now modestly positive ({pc(b('13+','v6'),2)}, interval [{s(b('13+','v6')['lo'],2)}, {s(b('13+','v6')['hi'],2)}]). TVT flat.")
sigtvt = [f"{a} {band_lbl(bb)} ({pt(b(bb, a, 'tvt'), 2)})" for bb in BANDS for a in ARMS if issig(b(bb, a, "tvt"))]
setp("4–6 loses on", f"**4–6** loses on v4 and v6 ({pc(b('4-6','v4'),2)}, {pc(b('4-6','v6'),2)}); v5b is {pt(b('4-6','v5b'),2)}. "
     f"**7–9** gains on v5b and v6 ({pc(b('7-9','v5b'),2)}, {pc(b('7-9','v6'),2)}); v4 {pt(b('7-9','v4'),2)}. **1–3** is flat on revenue (v5b impressions {pt(b('1-3','v5b','imps'),2)}). "
     + ("The only significant band-level TVT move is " + "; ".join(sigtvt) + "." if sigtvt else "No band-level TVT move is significant."))
setp("Device-level raw Welch, unbudgeted revenue/impressions, among devices who watched", "Device-level raw Welch, unbudgeted revenue/impressions, among devices who watched a targeted title of that type.")
bel = lambda a, bd_, t: R[f"below|{a}|{bd_}|{t}"][a]
below_rev = [bel(a, bd_, "below")["rev"]["l"] for a in ARMS for bd_ in ("10-12", "13+")]
above_rev = [bel(a, bd_, "at/above")["rev"]["l"] for a in ARMS for bd_ in ("10-12", "13+")]
tvt_sig = [f"{a} {bd_.replace('-', '–')} {t.replace('at/above', 'at/above control')} {pt(bel(a, bd_, t)['tvt'],2)}" for a in ARMS for bd_ in ("10-12", "13+") for t in ("below", "at/above") if issig(bel(a, bd_, t)["tvt"])]
setp("This reproduces the v8 defect", f"This reproduces the v8 defect on every arm. Titles served **below control** lose {abs(max(below_rev)):.1f}–{abs(min(below_rev)):.1f}% unbudgeted revenue; identical-band titles kept **at or above control** gain {min(above_rev):.1f}–{max(above_rev):.1f}%, "
     f"with little TVT give-back (exceptions: {'; '.join(tvt_sig)}). Direct support for a per-title floor: never serve fewer cue points than the ABF2 control. How much of each arm sits below control in the dense bands (§B.4) drives its band result in §2a.")
setp("Control n: movies", f"Control n: movies {n(mv['v4']['ctrl_n'])} · episodes {n(ep['v4']['ctrl_n'])}. Movies gain on all three arms (v4 {pc(mv['v4']['rev'],2)}, v5b {pc(mv['v5b']['rev'],2)}, v6 {pc(mv['v6']['rev'],2)} — biggest on v6, which was built for 60–120 min movies). "
     f"Episodes lose on v4 and v6 and are flat-to-negative on v5b ({pt(ep['v5b']['rev'],2)}), with no TVT give-back. The global unbudgeted picture is this mix — movie gains netting against episode losses.")
bp = {(x["band"], x["arm"]): x["sim_delta"] for x in G["band_pooled"]}
import numpy as np
keys = [(bb, a) for bb in BANDS for a in ARMS]
CORR = float(np.corrcoef([bp[k] for k in keys], [R["band|" + bb][a]["imps"]["l"] for bb, a in keys])[0, 1])
m = re.search(r"r = (0\.\d\d)", find("The simulation reproduces the measured pattern").text)
rep(f"r = {m.group(1)}", f"r = {CORR:.2f}")
f = {(x["tt"], x["arm"]): x for x in G["first"]}
fe = lambda a, k="first_srv": f"~{f[('Episode', a)][k]:.0f}s"
setp("First fireable break after the gate", "**First fireable break after the gate (TVT-weighted seconds):** "
     f"episodes ABF2 {fe('v4','first_ctl')}; v4 {fe('v4')}, v6 {fe('v6')} (slightly later), v5b {fe('v5b')} (earlier). "
     f"Movies ~{f[('Movie','v4')]['first_ctl']:.0f}s control; v4 {s(f[('Movie','v4')]['first_srv']-f[('Movie','v4')]['first_ctl'],0)}s, v6 {s(f[('Movie','v6')]['first_srv']-f[('Movie','v6')]['first_ctl'],0)}s, v5b {s(f[('Movie','v5b')]['first_srv']-f[('Movie','v5b')]['first_ctl'],0)}s vs control. "
     f"Unweighted title averages show a bigger episode gap (ABF2 {fe('v4','unw_first_ctl')}, v4 {fe('v4','unw_first_srv')}, v6 {fe('v6','unw_first_srv')}, v5b {fe('v5b','unw_first_srv')}), so the late first break sits mostly in lower-TVT titles.")
setp("Band and movie/episode numbers below", "Band and movie/episode numbers below are raw device-level Welch lifts (unbudgeted revenue/impressions, not CUPED), among devices who watched at least one targeted title in that band or type during the window. The same device can appear in more than one band. Titles are the 224,739 targeted titles; bands are the title’s ABF2 cue-point count.")
setp("Read this as an early read", "Read this as an early read on the extended window (through 2026-10-19), not a graduation call. Judge the money on total revenue, because the pruners add ad load that **budgeted/guaranteed** demand fills — so the open-auction pool alone understates the gain.")
setp("v5b is the leading revenue arm.", f"**v5b is the leading revenue arm.** Total revenue {pc(I['revt']['v5b'])}, budgeted {pc(I['revb']['v5b'])}, all-ad impressions {pc(I['impt']['v5b'])}, all significant. On the treated titles its unbudgeted revenue is also significantly positive ({pc(I['orevu']['v5b'])}). "
     f"The one caveat is a small TVT dip ({pc(I['tvt']['v5b'])} global, {pc(H['PV']['tvt']['v5b'])} on power viewers) — minor, but worth watching because v5b adds the most ad load.")
setp("v6 is the safer second.", f"**v6 is the safer second.** Total revenue {pc(I['revt']['v6'])}, budgeted {pc(I['revb']['v6'])}, all-ad impressions {pc(I['impt']['v6'])}, on-title unbudgeted {pc(I['orevu']['v6'])} — all significant, TVT flat ({pt(I['tvt']['v6'])}). "
     "A smaller, steadier gain than v5b with less added load and no engagement signal.")
setp("v4 is roughly parity", f"**v4 is roughly parity, not a clear loss.** Global total and unbudgeted revenue are both not significant ({pc(I['revt']['v4'])}, {pc(I['revu']['v4'])}), near its break-even design; but it is significantly negative on the treated titles (on-title unbudgeted {pc(I['orevu']['v4'])}). "
     f"Weakest of the three. TVT {pt(I['tvt']['v4'])}.")
setp("TVT is safe.", f"**TVT is safe.** The only significant TVT move is v5b’s {pc(I['tvt']['v5b'],2)} global ({pc(I['otvt']['v5b'],2)} on-title, {pc(H['PV']['tvt']['v5b'],2)} on power viewers); v4 and v6 are not significant. No arm trades meaningful engagement for the revenue. v5b’s small dip is the only item to monitor.")
ec = {x["band"]: x for x in G["cells"] if x["tt"] == "Episode" and x["arm"] == "v4"}
e13 = [cell[("Episode", "13+", a)] for a in ARMS]
setp("The episode loss is understood", f"**The episode loss is understood and fixable — it is upside, not a blocker.** v4 and v6 lose unbudgeted revenue on episodes (v5b {pt(ep['v5b']['rev'],2)}); the gate replay shows dense episodes fire fewer breaks ({s(max(e13),1)}% to {s(min(e13),1)}% in band 13+). "
     "For v4/v6 this is mostly a count miss (pruned below ABF2 on ~75% of dense episodes); for v5b mostly placement. A per-title floor plus an earlier episode cue would recover it at a TVT give-back bounded by the flat episode TVT here — so v5b’s and v6’s total-revenue gains should grow once episodes are fixed.")
setp("A longer window can still", "**A longer window can still move the arms.** v5b and v6 are already significant on total revenue; v4 is parity. The extension from 10-15 to 10-19 removed v4’s small TVT “gain” and, with the corrected title set, the on-title gains are larger than in earlier versions; direction is stable across populations.")

# B.5: refreshed pool table + stale-note
hp = find("B.5 Revenue-pool diagnostic")
insert_after(hp, "**Window note (2026-10-20):** the impression-pool table below is refreshed on the 2026-10-09 → 2026-10-19 window. The pacing, shift-share, rollout-range and pre-roll analysis in this section "
             "(the verdict and items 1–6, with their device counts, dollar figures and percentages) was computed on the earlier 2026-10-09 → 2026-10-15 window and was not re-run — read it as the 10-15 analysis.",
             template=find("The gain is concentrated in budgeted"))

# Appendix C
setp("Window. Outcomes", "**Window.** Outcomes 2026-10-09 through 2026-10-19, capped at 2026-10-20 00:00 UTC (the exposures table is loaded through the end of 2026-10-19). CUPED covariate 2026-09-25 → 2026-10-08, before any v10 exposure. "
     "Active = ≥1 video session since first exposure in-window. ToT = ≥1 second of TVT on a targeted title (the 224,739 titles below) in-window.")
setp("ABF2 title list.", "**Title set.** On-title, ToT, band, type and below-control tables use the 224,739 titles in `ml_dev.tmp.abf3_v10_targeted_titles` (191,035 episodes + 33,704 movies): the ABF2-catalog titles that were served v10 cue points. "
     "The ABF2 catalog itself is movie/episode streams in `core_prod.tubidw` with non-empty `cue_points_abf`, excluding `cue_points_info = 'abf3_pruned_v8'` and `core_dev.dsa.temp_bp_abf_v8_all_titles` (299,272 live titles); the other ~74.5k catalog titles were not served v10 cue points, carry identical cues in every arm, and are excluded. "
     "Versions of this report before 2026-10-20 used the whole catalog and so understated every on-title and band lift by roughly a quarter. Bands use the title’s ABF2 control cue-point count from the live stream tables. The CUPED covariate for on-title outcomes is the all-title pre-period, matching the v8 notebook.")
setp("Reconciliation with Statsig.", "**Reconciliation with Statsig.** Done on the 2026-10-15 window: total, all-ad impressions, and unbudgeted agreed closely with the scorecard; budgeted read a little higher here (v5b +2.39% vs ~+1.77% on the tile) — expected from CUPED vs raw and a slightly different budgeted tagging; "
     "budgeted is also the most window-sensitive metric. Not re-run for the extended window (no scorecard pull in this refresh).")
sp = find("Served cue lists.")
if sp is not None:
    txt = sp.text
    set_markup(sp, "**Served cue lists.**" + txt[len("Served cue lists."):] + " The table now also holds extra arms that are not part of v10; they are excluded everywhere.")
vp = find("Verification (2026-10-15).")
set_markup(vp, "**Verification (2026-10-20).** Refresh on the extended window with the corrected title set. All Databricks queries were re-run (headline CUPED, band/type/below-control, gate replay, viewing depth). "
           "CUPED and Welch are computed inside SQL; the SQL estimator was checked against the Python estimator on the OTT ITT total-revenue sums (match to 0.001pp). "
           "The title-set correction was checked by running the same estimator on an independent OTT analysis’s window and pre-period (10-10 → 10-18; 09-17 → 09-30): global unbudgeted revenue and TVT and the targeted-title on-title lifts matched it to rounding (e.g. on-title v5b +0.90%), while the 299k-catalog definition did not (+0.67%). "
           "The gate replay was ported from the notebook to SQL and validated by reproducing the 10-15 OTT cells, first-break timings and episode decomposition to rounding (OTT keeps all-platform in-window TVT weights); it already ran on served titles only, so the title-set correction does not change it. "
           "Not refreshed: the §B.5 pacing / pre-roll analysis (10-15 window).")

json.dump(MISSES, open(OUT + ".misses.json", "w"))
d.save(OUT)
print("saved", OUT, "misses:", MISSES)
