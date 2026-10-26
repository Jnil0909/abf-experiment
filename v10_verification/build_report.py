"""Build an ABF3 v10 results doc (Mobile / Web) in the same format as the OTT report.

Usage: build_report.py <template.docx> <values.json> <gate_base.json> <gate_tol.json> <served.json> <narrative.json> <out.docx>
Tables come from the Databricks outputs; prose comes from narrative.json (written after reading the results).
"""
import json, sys, docx
from docx.text.paragraph import Paragraph

TPL, VALS, GB, GT, SRV, NAR, OUT = sys.argv[1:8]
V = json.load(open(VALS)); H = V["head"]; R = V["raw"]
GBASE = json.load(open(GB)); GTOL = json.load(open(GT))
SV = json.load(open(SRV)); N = json.load(open(NAR))
ARMS = ("v4", "v5b", "v6")
BANDS = ("1-3", "4-6", "7-9", "10-12", "13+")
TABLE_STYLE = "Light Grid Accent 1"

d = docx.Document(TPL)
body = d.element.body
for el in list(body):
    if not el.tag.endswith("sectPr"):
        body.remove(el)


# ---------------- formatting helpers ----------------
def s(x, nd):
    x = round(x, nd) + 0.0
    return f"{x:+.{nd}f}".replace("-", "−")


def sig(e):
    return "sig" if e["lo"] > 0 or e["hi"] < 0 else "ns"


def full(e, nd=3, tag=True):
    t = f"{s(e['l'], nd)}% [{s(e['lo'], nd)}, {s(e['hi'], nd)}]"
    return t + (" " + sig(e) if tag else "")


def short(e, nd=2):
    return f"{s(e['l'], nd)}% {sig(e)}"


def tvt_cell(e, nd=2):
    return full(e, nd) if sig(e) == "sig" else short(e, nd)


n = lambda x: f"{x:,}"
band_lbl = lambda b: b.replace("-", "–")


def para(text="", lead=None, style="Normal"):
    p = d.add_paragraph(style=style)
    if lead:
        p.add_run(lead).bold = True
        p.add_run(" " + text if text else "")
    else:
        p.add_run(text)
    return p


def heading(text, level=1):
    return d.add_paragraph(text, style=f"Heading {level}")


def bullets(items):
    for it in items:
        if isinstance(it, list):
            para(it[1], lead=it[0], style="List Bullet")
        else:
            para(it, style="List Bullet")


def table(header, rows):
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    t.style = d.styles[TABLE_STYLE]
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for ri, r in enumerate(rows, start=1):
        for ci, v in enumerate(r):
            t.rows[ri].cells[ci].text = str(v)
    d.add_paragraph()
    return t


# ---------------- document ----------------
d.add_paragraph(N["title"], style="Title")
for lead, txt in N["front"]:
    para(txt, lead=lead)

heading("TL;DR")
bullets(N["tldr"])
I = H["ITT"]
table(["Arm", "Total revenue", "Budgeted revenue", "Unbudgeted revenue", "All-ad impressions", "TVT"],
      [[N["arm_label"][g], full(I["revt"][g]), short(I["revb"][g], 3), short(I["revu"][g], 3),
        short(I["impt"][g], 3), short(I["tvt"][g], 3)] for g in ARMS])
para(N["tldr_note"])

heading(f"1. Headline ({N['window_short']})")
para(N["headline_intro"])
para("Revenue, three ways (global ITT — the headline money table). Total = all billable ads; budgeted = campaigns carrying a budget (direct/guaranteed/PMP); unbudgeted = open-auction only.")
table(["Arm", "Total revenue", "Budgeted revenue", "Unbudgeted revenue", "All-ad impressions"],
      [[g, full(I["revt"][g]), full(I["revb"][g]), full(I["revu"][g]), full(I["impt"][g])] for g in ARMS])
cm = I["ctrl_mean"]
para(f"Control per-device means: total revenue {cm['revt']:.3f}, budgeted {cm['revb']:.3f}, all-ad impressions {cm['impt']:.1f}; "
     f"ITT n/arm ≈ {H['ITT']['n']['Control']/1e6:.2f}M. " + N["money_read"])
para("Global unbudgeted detail (the v8-definition slice — all titles):")
rows = []
for pop in ("ITT", "Active", "ToT"):
    for g in ARMS:
        P = H[pop]
        rows.append([pop, g, n(P["n"][g]), full(P["revu"][g]), full(P["impu"][g]), full(P["tvt"][g])])
table(["Population", "Arm", "n", "Unbudgeted revenue", "Impressions", "TVT"], rows)
para(f"Control n: ITT {n(H['ITT']['n']['Control'])} · Active {n(H['Active']['n']['Control'])} · ToT {n(H['ToT']['n']['Control'])}.")
para("On-title (the 224,739 targeted titles — ABF2 titles served v10 cue points; same device counts as the matching global row):")
rows = []
for pop in ("ITT", "Active", "ToT"):
    for g in ARMS:
        P = H[pop]
        rows.append([pop, g, full(P["orevu"][g]), full(P["oimpu"][g]), full(P["otvt"][g])])
table(["Population", "Arm", "Unbudgeted revenue", "Impressions", "TVT"], rows)
para(N["ontitle_read"])
pvs = H["PV"]["n"]["Control"] / H["ITT"]["n"]["Control"] * 100
para("Power viewer = ≥8 visit-days in last 28 days and >4 watch-hours/visit-day from the latest `device_okr_daily` snapshot "
     f"on/before first exposure. Everyone else (incl. devices absent from the table) is non-power. Flag set pre-exposure; power share ~{pvs:.1f}% of Control.",
     lead="NPV_TVT — TVT split by viewer type before first exposure.")
rows = []
for pop, lbl in (("NPV", "Non-power viewers"), ("PV", "Power viewers")):
    for g in ARMS:
        e = H[pop]["tvt"][g]
        rows.append([lbl, g, n(H[pop]["n"][g]), full(e, 3, tag=False), "Yes" if sig(e) == "sig" else "No"])
table(["Subpopulation", "Arm", "n", "TVT (CUPED)", "Stat-sig"], rows)
para(f"Control n: non-power {n(H['NPV']['n']['Control'])} · power {n(H['PV']['n']['Control'])}. " + N["npv_read"])

heading("2. Where the result lives")
para("Band and movie/episode numbers below are raw device-level Welch lifts (unbudgeted revenue/impressions, not CUPED), among devices who "
     "watched at least one targeted title in that band or type during the window. The same device can appear in more than one band. "
     "Titles are the 224,739 targeted titles; bands are the title’s ABF2 cue-point count.")
heading("2a. By ABF2 density band", 2)
rows = []
for b in BANDS:
    for g in ARMS:
        x = R["band|" + b][g]
        rows.append([band_lbl(b), g, full(x["rev"], 2), short(x["tvt"]), full(x["imps"], 2), n(x["n"])])
table(["ABF2 band", "Arm", "Unbudgeted revenue", "TVT", "Impressions", "n (devices)"], rows)
para("Control n by band: " + " · ".join(f"{band_lbl(b)} {n(R['band|'+b]['v4']['ctrl_n'])}" for b in BANDS) + ".")
bullets(N["band_read"])

heading("2b. Below-control versus at-or-above-control", 2)
para("Built from the served cue lists in `ml_dev.tmp.abf3_new_pruner_cuepoints` (snapshot 2026-10-08; T1 = v4, T2 = v6, T3 = v5b). "
     "Each dense-band title is tagged by whether that arm kept fewer cue points than its ABF2 control count or at or above it. "
     "Raw device-level Welch, treatment versus control, on those titles.")
rows = []
for b in ("10-12", "13+"):
    for tag, lbl in (("below", "below control"), ("at/above", "at/above control")):
        for g in ARMS:
            x = R[f"below|{g}|{b}|{tag}"][g]
            rows.append([band_lbl(b), lbl, g, full(x["rev"], 2), tvt_cell(x["tvt"])])
table(["Band", "Served vs control", "Arm", "Unbudgeted revenue", "TVT"], rows)
para(N["below_read"])

heading("2c. Movies versus episodes", 2)
para("Device-level raw Welch, unbudgeted revenue/impressions, among devices who watched a targeted title of that type.")
rows = []
for tt in ("Movie", "Episode"):
    for g in ARMS:
        x = R["type|" + tt][g]
        rows.append([tt, g, full(x["rev"], 2), short(x["tvt"]), short(x["imps"]), n(x["n"])])
table(["Type", "Arm", "Unbudgeted revenue", "TVT", "Impressions", "n (devices)"], rows)
para(f"Control n: movies {n(R['type|Movie']['v4']['ctrl_n'])} · episodes {n(R['type|Episode']['v4']['ctrl_n'])}. " + N["type_read"])

cp = {(r[1], r[2]): r for r in SV["rows"] if r[0] == "cphr"}
para("Average first cue point (title-level, served lists — the same catalog-level lists serve every platform):",
     lead="Placement — first-cue timing from the served lists.")
mins = lambda x: f"~{float(x):.0f}s ({float(x)/60:.1f} min)"
table(["Arm", "First cue, episodes", "First cue, movies"],
      [["ABF2 control", mins(cp[("Episode", "v4")][8]), mins(cp[("Movie", "v4")][8])]] +
      [[g, mins(cp[("Episode", g)][9]), mins(cp[("Movie", g)][9])] for g in ARMS])

def gate_tbl(G):
    c = {(x["tt"], x["band"], x["arm"]): x["sim_delta"] for x in G["cells"]}
    return [[tt, band_lbl(b)] + [f"{s(c[(tt, b, g)], 1)}%" for g in ARMS]
            for tt in ("Movie", "Episode") for b in BANDS]

para(N["gate_intro"], lead="Gate replay — fireable breaks from the served positions.")
para("Profile A — OTT base gates (episodes 420s initial / 600s spacing, movies 900s / 720s). TVT-weighted simulated fired-break delta, arm vs ABF2:")
table(["Type", "Band", "v4", "v5b", "v6"], gate_tbl(GBASE))
para("Profile B — tolerant gates (episodes 360s / 360s, movies 540s / 420s), closer to the earlier break starts observed on this platform:")
table(["Type", "Band", "v4", "v5b", "v6"], gate_tbl(GTOL))
para(N["gate_read"])
para("Episode deficit — count vs position (TVT-weighted, pooled over episodes; total = fewer cue points + placement + gains on titles kept at/above control):")
rows = []
for lbl, G in (("A (OTT base)", GBASE), ("B (tolerant)", GTOL)):
    for x in G["episode_decomp"]:
        rows.append([lbl, x["arm"], f"{s(x['total'],1)}%", f"{s(x['from_fewer'],1)}%", f"{s(x['from_placement'],1)}%", f"{s(x['from_gains'],1)}%"])
table(["Gate profile", "Arm", "Total episode simΔ", "from fewer cue points", "from placement", "from gains"], rows)
para(N["decomp_read"])

heading("3. Decision frame")
para(N["decision_intro"])
bullets(N["decision"])

heading("Appendix B — served cue points")
heading("B.4 Cue points per hour versus ABF2", 2)
para("From the served counts and durations in `ml_dev.tmp.abf3_new_pruner_cuepoints` (catalog-level; identical on every platform). Title-average cue points per hour vs the ABF2 control:")
def cph(tt):
    base = float(cp[(tt, "v4")][6])
    return [tt, f"{base:.2f}"] + [f"{float(cp[(tt, g)][7]):.2f} ({s((float(cp[(tt, g)][7])/base-1)*100, 0)}%)" for g in ARMS]
table(["Type", "ABF2", "v4", "v5b", "v6"], [cph("Episode"), cph("Movie")])
bd = {(r[1], r[4], r[2]): r for r in SV["rows"] if r[0] == "band"}
rows = []
for tt, bands in (("Episode", ("7-9", "10-12", "13+")), ("Movie", ("10-12", "13+"))):
    for b in bands:
        for g in ARMS:
            r = bd[(tt, b, g)]
            rows.append([tt, band_lbl(b), g, r[10], r[11], r[12]])
table(["Type", "Band", "Arm", "% below", "% equal", "% more"], rows)
heading("B.5 Revenue-pool diagnostic — is the budgeted lift incremental?", 2)
table(["Arm", "Budgeted impressions", "Unbudgeted impressions"],
      [[g, short(I["impb"][g], 3), short(I["impu"][g], 3)] for g in ARMS])
para(N["pool_read"])

heading("Appendix C — Methods")
bullets(N["methods"])

d.save(OUT)
print("saved", OUT)
